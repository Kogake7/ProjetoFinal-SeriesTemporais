"""T01: fontes originais dos grupos -> tratamento causal -> cinco Excel comuns.

Execute ``python preparacao_bases.py`` a partir de qualquer diretório.
Os dados originais permanecem nas pastas dos grupos. A saída contém um Excel
tratado por grupo, em ``dados_tratados/``.
"""
from __future__ import annotations

import argparse
import json
import os
from pathlib import Path
import tempfile

import numpy as np
import pandas as pd

from validacao_bases import (
    CONFIGURACOES, RAIZ_PROJETO, calcular_sha256, carregar_base,
    extrair_datas, validar_base,
)

PROTOCOLO = {
    'bitcoin': dict(grupo=1, frequencia='D', seed=42, treino=0.70, alvo='priceClose', limite_ffill=1),
    'trafego': dict(grupo=2, frequencia='h', seed=67, treino=0.80, alvo='traffic_volume', limite_ffill=6),
    'poluicao': dict(grupo=3, frequencia='h', seed=42, treino=0.80, alvo='PM2.5', limite_ffill=6),
    'clima': dict(grupo=4, frequencia='h', seed=42, treino=0.80, alvo='T (degC)', limite_ffill=6),
    'ouro': dict(grupo=5, frequencia='W-FRI', seed=42, treino=0.75, alvo='VALUE', limite_ffill=0),
}
CONTROLES = ['_alvo_observado', '_linha_completa', '_particao']


def tratar_base(nome: str, bruto: pd.DataFrame) -> tuple[pd.DataFrame, dict]:
    """Tratamento determinístico; sem escala, interpolação futura ou alvo inventado.

    Retorna grade regular completa. Nulos remanescentes são ausências reais:
    os modelos devem criar lags ANTES de selecionar exemplos completos.
    """
    cfg = PROTOCOLO[nome]
    fonte = CONFIGURACOES[nome]
    df = bruto.copy()
    rel = dict(linhas_brutas=len(df), duplicatas_exatas=int(df.duplicated().sum()))
    df = df.drop_duplicates()
    datas = extrair_datas(df, fonte)
    if getattr(datas.dt, 'tz', None) is not None:
        datas = datas.dt.tz_convert('UTC').dt.tz_localize(None)
    rel['datas_invalidas'] = int(datas.isna().sum())
    df.index = pd.DatetimeIndex(datas, name='data')
    df = df.loc[df.index.notna()].sort_index(kind='stable')
    remover = list(fonte.colunas_data) + ([fonte.coluna_data] if fonte.coluna_data else [])
    df = df.drop(columns=remover + (['No', 'station'] if nome == 'poluicao' else []), errors='ignore')
    if nome == 'bitcoin':
        # Epochs de fechamento/máximas não são medidas úteis para esta interface.
        df = df.drop(columns=['timeClose', 'timeHigh', 'timeLow'], errors='ignore')
    if nome == 'trafego':
        df['holiday'] = df['holiday'].fillna('No Holiday')
    categoricas = [c for c in ['holiday', 'weather_main', 'weather_description', 'wd'] if c in df]
    numericas = df.columns.difference(categoricas)
    df[numericas] = df[numericas].apply(pd.to_numeric, errors='coerce')
    rel['sentinelas_menos_9999'] = int((df[numericas] == -9999).sum().sum())
    df[numericas] = df[numericas].replace([-9999, np.inf, -np.inf], np.nan)
    if nome == 'trafego':
        df.loc[df['temp'] <= 0, 'temp'] = np.nan  # Kelvin fisicamente inválido.
    nao_negativas = {
        'bitcoin': ['priceOpen', 'priceHigh', 'priceLow', 'priceClose', 'volume'],
        'trafego': ['traffic_volume', 'rain_1h', 'snow_1h'],
        'poluicao': ['PM2.5', 'PM10', 'SO2', 'NO2', 'CO', 'O3', 'RAIN', 'WSPM'],
        'clima': ['wv (m/s)', 'max. wv (m/s)'], 'ouro': ['VALUE'],
    }[nome]
    for col in nao_negativas:
        df.loc[df[col] < 0, col] = np.nan
    conflitos = df.groupby(level=0)[cfg['alvo']].nunique().gt(1)
    rel['timestamps_com_alvos_conflitantes'] = int(conflitos.sum())
    rel['timestamps_duplicados'] = int(df.index.duplicated().sum())
    # Mantém a primeira linha inteira (ordem estável), sem somar tráfego repetido.
    # Um alvo conflitante vira ausência, em vez de escolher uma verdade arbitrária.
    df = df.loc[~df.index.duplicated(keep='first')].copy()
    df.loc[conflitos[conflitos].index, cfg['alvo']] = np.nan
    if nome == 'clima':
        angulo = np.deg2rad(df['wd (deg)'])
        df['wd_sin'] = np.sin(angulo)
        df['wd_cos'] = np.cos(angulo)
        regras = {col: 'mean' for col in df.columns}
        regras['max. wv (m/s)'] = 'max'
        df = df.resample('h', closed='left', label='left').agg(regras)
        df['wd (deg)'] = np.rad2deg(np.arctan2(df['wd_sin'], df['wd_cos'])) % 360
        df.loc[np.hypot(df['wd_sin'], df['wd_cos']) < 1e-10, 'wd (deg)'] = np.nan
    elif nome == 'ouro':
        df = df.resample('W-FRI').last()  # Último preço válido da semana.
    else:
        df = df.asfreq(cfg['frequencia'])
    antes = df.isna().sum()
    exogenas = df.columns.difference([cfg['alvo']])
    if cfg['limite_ffill']:
        df[exogenas] = df[exogenas].ffill(limit=cfg['limite_ffill'])
    rel['preenchimentos_causais_por_coluna'] = (antes - df.isna().sum()).astype(int).to_dict()
    rel['nulos_restantes_por_coluna'] = df.isna().sum().astype(int).to_dict()
    rel['linhas_grade'] = len(df)
    rel['linhas_completas'] = int(df.notna().all(axis=1).sum())
    rel['alvos_observados'] = int(df[cfg['alvo']].notna().sum())
    validar_tratada(df, nome)
    corte = int(len(df) * cfg['treino'])
    meta = dict(base=nome, **cfg, inicio=str(df.index.min()), fim=str(df.index.max()),
                fim_treino=str(df.index[corte - 1]), inicio_teste=str(df.index[corte]),
                linhas_treino=corte, linhas_teste=len(df) - corte,
                fuso='UTC sem timezone no Excel' if nome == 'bitcoin' else 'horário local da fonte',
                diagnostico=rel)
    return df, meta


def validar_tratada(df: pd.DataFrame, nome: str) -> None:
    cfg = PROTOCOLO[nome]
    if len(df) < 2 or cfg['alvo'] not in df or df[cfg['alvo']].notna().sum() < 2:
        raise ValueError(f'{nome}: série insuficiente ou alvo ausente.')
    esperado = pd.date_range(df.index.min(), df.index.max(), freq=cfg['frequencia'], name='data')
    if not df.index.equals(esperado):
        raise ValueError(f'{nome}: datas inválidas, duplicadas, desordenadas ou fora da grade.')
    if np.isinf(df.select_dtypes(include='number').to_numpy()).any():
        raise ValueError(f'{nome}: infinito na base tratada.')


def carregar_base_tratada(nome: str, raiz: str | Path = RAIZ_PROJETO):
    """Lê o único Excel tratado do grupo e confere hash, grade e configuração.

    O DataFrame retornado exclui colunas de controle para evitar usá-las como features.
    O corte em metadados deve ser mantido depois da criação de lags.
    """
    pasta = Path(raiz) / 'dados_tratados'
    manifesto = json.loads((pasta / 'manifesto.json').read_text(encoding='utf-8'))
    item = manifesto['bases'][nome]
    caminho = pasta / item['excel']
    if calcular_sha256(caminho) != item['sha256_excel']:
        raise ValueError(f'Excel alterado: {caminho}. Regenere as bases tratadas.')
    df = pd.read_excel(caminho, sheet_name='dados', index_col='data', parse_dates=['data'])
    df = df.drop(columns=CONTROLES)
    # Restaura frequência sem remover os períodos ausentes.
    validar_tratada(df, nome)
    df = df.asfreq(item['frequencia'])
    for chave in ['grupo', 'frequencia', 'seed', 'treino', 'alvo']:
        if item[chave] != PROTOCOLO[nome][chave]:
            raise ValueError(f'Protocolo mudou ({chave}); regenere as bases tratadas.')
    return df, item


def exportar_excel(df: pd.DataFrame, meta: dict, destino: Path) -> None:
    """Exportação streaming em Python, com datas e números tipados.

    Fallback: o exportador Artifact Tool encerrou com falha nativa 0xC0000409
    neste Windows. openpyxl também permite regenerar os dados fora do Codex.
    """
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook(write_only=True)
    dados = wb.create_sheet('dados')
    metadados = wb.create_sheet('metadados')
    for sheet in [dados, metadados]:
        sheet.freeze_panes = 'B2' if sheet.title == 'dados' else 'A2'
    for i in range(1, len(df.columns) + 2):
        dados.column_dimensions[get_column_letter(i)].width = 24
    metadados.column_dimensions['A'].width = 36
    metadados.column_dimensions['B'].width = 105

    def cabecalho(sheet, valores):
        cells = []
        for valor in valores:
            cell = WriteOnlyCell(sheet, valor)
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='16324F')
            cells.append(cell)
        sheet.append(cells)

    cabecalho(dados, ['data', *df.columns])
    for linha in df.itertuples(index=True, name=None):
        data = WriteOnlyCell(dados, linha[0].to_pydatetime())
        data.number_format = 'yyyy-mm-dd hh:mm'
        dados.append([data, *(None if pd.isna(v) else v for v in linha[1:])])
    dados.auto_filter.ref = f'A1:{get_column_letter(len(df.columns) + 1)}{len(df) + 1}'
    cabecalho(metadados, ['campo', 'valor'])
    def campos(valores, prefixo=''):
        for chave, valor in valores.items():
            nome = f'{prefixo}{chave}'
            if isinstance(valor, dict):
                yield from campos(valor, nome + '.')
            else:
                yield nome, json.dumps(valor, ensure_ascii=False) if isinstance(valor, list) else valor

    for i, (chave, valor) in enumerate(campos(meta), start=2):
        cells = [WriteOnlyCell(metadados, chave), WriteOnlyCell(metadados, valor)]
        for cell in cells:
            cell.alignment = Alignment(wrap_text=True, vertical='top')
        metadados.row_dimensions[i].height = max(30, 15 * (len(str(valor)) // 100 + 1), 15 * (len(chave) // 34 + 1))
        metadados.append(cells)
    wb.save(destino)


def preparar_todas(raiz: str | Path = RAIZ_PROJETO, *, sobrescrever: bool = False) -> Path:
    """Gera exatamente cinco Excel tratados a partir das fontes em grupo1...5.

    Regerar exige ``sobrescrever=True`` e substitui os mesmos arquivos.
    O manifesto registra hashes e regras, mas não é uma sexta base.
    """
    raiz = Path(raiz).resolve()
    destino = raiz / 'dados_tratados'
    esperados = [destino / f'base-tratada-{cfg["grupo"]}.xlsx'
                 for cfg in PROTOCOLO.values()]
    if not sobrescrever and any(p.exists() for p in esperados):
        raise FileExistsError('Já existem bases tratadas. Use --sobrescrever para regenerar as mesmas cinco.')
    destino.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix='.preparacao-', dir=destino) as tmp:
        trabalho = Path(tmp)
        manifesto = dict(bases={})
        for nome in PROTOCOLO:
            print(f'Tratando {nome}...', flush=True)
            df, meta = tratar_base(nome, carregar_base(nome, raiz))
            meta['diagnostico_bruto'] = validar_base(nome, raiz).to_dict()
            meta['fonte'] = CONFIGURACOES[nome].caminho
            meta['sha256_fonte'] = calcular_sha256(raiz / meta['fonte'])
            meta['procedencia'] = ('Arquivo local adaptado; link oficial pendente de fornecimento.'
                                  if nome == 'ouro' else 'Arquivo original preservado na pasta do grupo.')
            meta['politica'] = ('Grade regular; alvo nunca imputado; exógenas com ffill limitado; '
                               'lags antes de filtrar linhas; escala e encoding ajustados no treino. '
                               'Outliers plausíveis preservados para análise T03.')
            tabela = df.copy()
            tabela['_alvo_observado'] = df[meta['alvo']].notna()
            tabela['_linha_completa'] = df.notna().all(axis=1)
            tabela['_particao'] = np.where(df.index < pd.Timestamp(meta['inicio_teste']), 'treino', 'teste')
            arquivo = f'base-tratada-{meta["grupo"]}.xlsx'
            exportar_excel(tabela, meta, trabalho / arquivo)
            meta['excel'] = arquivo
            meta['sha256_excel'] = calcular_sha256(trabalho / arquivo)
            manifesto['bases'][nome] = meta
        (trabalho / 'manifesto.json').write_text(json.dumps(manifesto, ensure_ascii=False, indent=2), encoding='utf-8')
        # Só altera a saída após gerar as cinco bases com sucesso.
        for meta in manifesto['bases'].values():
            os.replace(trabalho / meta['excel'], destino / meta['excel'])
        os.replace(trabalho / 'manifesto.json', destino / 'manifesto.json')
    return destino


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--sobrescrever', action='store_true',
                        help='Regenera os mesmos cinco Excel tratados.')
    args = parser.parse_args()
    print(preparar_todas(sobrescrever=args.sobrescrever))
